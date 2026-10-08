"""Build the demo slides (.pptx) and the speaking script (.md) from slides/figures/.
Dựng slide demo (.pptx) và kịch bản thuyết trình (.md) từ slides/figures/.

Run scripts/export_figures.py first, then: python slides/build_slides.py
Chạy scripts/export_figures.py trước, rồi: python slides/build_slides.py
"""

import json
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
FIG = HERE / "figures"
OUT_PPTX = HERE / "demo_slides.pptx"
OUT_MD = HERE / "kich_ban_demo.md"

FONT = "Arial"
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
TEXT = RGBColor(0x33, 0x33, 0x33)
GREY = RGBColor(0x7F, 0x8C, 0x8D)
LIGHT = RGBColor(0xF2, 0xF4, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RED = RGBColor(0xC0, 0x39, 0x2B)
BLUE = RGBColor(0x29, 0x80, 0xB9)
GREEN = RGBColor(0x27, 0xAE, 0x60)
AMBER = RGBColor(0xFD, 0xF2, 0xE0)

REFERENCES = [
    "Kim, D. & Solomon, M. G. (2018). Fundamentals of Information Systems Security (3rd ed.), Ch. 5. Jones & Bartlett Learning.",
    "Schroff, F., Kalenichenko, D. & Philbin, J. (2015). FaceNet: A unified embedding for face recognition and clustering. CVPR.",
    "Serengil, S. I. & Ozpinar, A. (2020). LightFace: A hybrid deep face recognition framework. ASYU, 23–27.",
    "Teoh, A. B. J., Ngo, D. C. L. & Goh, A. (2004). Biohashing: two factor authentication featuring fingerprint data "
    "and tokenised random number. Pattern Recognition 37(11), 2245–2255.",
    "Hahn, V. K. & Marcel, S. (2022). Towards protecting face embeddings in mobile face verification scenarios. "
    "IEEE T-BIOM 4, 117–134.",
    "Adler, A. (2003). Sample images can be independently restored from face recognition templates. CCECE, vol. 2, 1163–1166.",
    "Adler, A. (2004). Images can be regenerated from quantized biometric match score data. CCECE.",
]


def vn(x: float, d: int = 2) -> str:
    """Vietnamese number format: decimal comma, true minus sign.
    Định dạng số kiểu Việt: dấu phẩy thập phân, dấu trừ chuẩn.
    """
    return f"{x:.{d}f}".replace(".", ",").replace("-", "−")


# ---------- Drawing helpers / Hàm vẽ ----------


def _add_runs(paragraph, text: str, size: float, color: RGBColor, bold: bool = False, font: str = FONT) -> None:
    """Add text where **...** marks bold segments.
    Thêm chữ, đoạn nằm giữa **...** được in đậm.
    """
    for i, part in enumerate(text.split("**")):
        if not part:
            continue
        run = paragraph.add_run()
        run.text = part
        run.font.size = Pt(size)
        run.font.name = font
        run.font.color.rgb = color
        run.font.bold = bold or (i % 2 == 1)


def textbox(slide, x, y, w, h, lines, size=16, color=TEXT, bold=False, align=PP_ALIGN.LEFT,
            anchor=MSO_ANCHOR.TOP, bullets=False, space_after=6, font=FONT):
    """Text box with one paragraph per line; optional real bullets with hanging indent.
    Hộp chữ, mỗi dòng một đoạn; có thể dùng bullet thật với thụt lề treo.
    """
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, Inches(0.04))
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        _add_runs(p, line, size, color, bold, font)
        if bullets:
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(int(Inches(0.27))))
            pPr.set("indent", str(-int(Inches(0.27))))
            etree.SubElement(pPr, qn("a:buChar")).set("char", "•")
    return box


def panel(slide, x, y, w, h, fill=LIGHT, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    shp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(1.5)
    shp.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        shp.adjustments[0] = 0.08
    return shp


def label_box(slide, x, y, w, h, text, fill, color=WHITE, size=14, bold=True):
    """Filled box with centered text, used in the pipeline diagram.
    Hộp tô màu có chữ ở giữa, dùng trong sơ đồ luồng xử lý.
    """
    shp = panel(slide, x, y, w, h, fill=fill)
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, Inches(0.05))
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        _add_runs(p, line, size if i == 0 else size - 2, color, bold if i == 0 else False)
    return shp


def arrow(slide, x1, y1, x2, y2, color=GREY):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    conn.line.color.rgb = color
    conn.line.width = Pt(2)
    ln = conn.line._get_or_add_ln()
    etree.SubElement(ln, qn("a:tailEnd")).set("type", "triangle")


def picture(slide, path: Path, x, y, max_w, max_h):
    """Insert an image scaled to fit the box, centered.
    Chèn ảnh co vừa khung, căn giữa.
    """
    with Image.open(path) as im:
        ratio = im.width / im.height
    w, h = (max_w, max_w / ratio) if max_w / ratio <= max_h else (max_h * ratio, max_h)
    slide.shapes.add_picture(str(path), Inches(x + (max_w - w) / 2), Inches(y + (max_h - h) / 2), Inches(w), Inches(h))


def table(slide, x, y, w, rows, col_widths, size=13, row_h=0.42, header_fill=NAVY):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(row_h * len(rows)))
    tbl = shape.table
    for j, cw in enumerate(col_widths):
        tbl.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(row_h)
        for j, value in enumerate(row):
            cell = tbl.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = header_fill if i == 0 else (LIGHT if i % 2 else WHITE)
            cell.margin_left = cell.margin_right = Inches(0.06)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
            _add_runs(p, str(value), size, WHITE if i == 0 else TEXT, bold=(i == 0))
    return tbl


def header(slide, title: str, tag: str | None = None) -> None:
    textbox(slide, 0.5, 0.3, 10.2, 0.8, [title], size=27, color=NAVY, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    panel(slide, 0.5, 1.12, 1.3, 0.07, fill=RED, shape=MSO_SHAPE.RECTANGLE)
    if tag:
        label_box(slide, 10.75, 0.45, 2.1, 0.48, tag, fill=LIGHT, color=NAVY, size=12, bold=True)


def footer(slide, source: str, number: int) -> None:
    textbox(slide, 0.5, 6.98, 11.4, 0.35, [source], size=10, color=GREY, anchor=MSO_ANCHOR.MIDDLE)
    textbox(slide, 12.2, 6.98, 0.65, 0.35, [str(number)], size=10, color=GREY, align=PP_ALIGN.RIGHT,
            anchor=MSO_ANCHOR.MIDDLE)


def formula(slide, x, y, w, h, segments, size=20, color=NAVY):
    """Centered one-line formula; each segment is (text, baseline) with baseline "sub", "sup" or None.
    Công thức một dòng căn giữa; mỗi đoạn là (chữ, vị trí) với vị trí "sub", "sup" hoặc None.
    """
    box = textbox(slide, x, y, w, h, [""], size=size, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    p = box.text_frame.paragraphs[0]
    for text, baseline in segments:
        run = p.add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.name = FONT
        run.font.color.rgb = color
        run.font.bold = True
        # Baseline offset in thousandths of a percent / Độ lệch baseline theo phần nghìn phần trăm
        if baseline:
            run.font._rPr.set("baseline", "-25000" if baseline == "sub" else "30000")
    return box


def callout(slide, x, y, w, h, text, fill=AMBER, size=15):
    shp = panel(slide, x, y, w, h, fill=fill)
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for side in ("margin_left", "margin_right"):
        setattr(tf, side, Inches(0.15))
    _add_runs(tf.paragraphs[0], text, size, TEXT)
    return shp


# ---------- Slides / Các slide ----------


def build(n: dict) -> tuple[Presentation, list[dict]]:
    """Create all slides; also return per-slide script entries for the markdown file.
    Tạo toàn bộ slide; trả kèm lời thoại từng slide cho file markdown.
    """
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    script: list[dict] = []

    t1, t2, t3, t4 = n["tab1"], n["tab2"], n["tab3"], n["tab4"]
    th = vn(n["threshold"])
    crossed = [r["crossed_at"] for r in t2["runs"] if r["crossed_at"] is not None]
    finals = [r["final"] for r in t2["runs"]]
    pair = " và ".join(n["example_pair"])

    # 1. Title / Trang tiêu đề
    s = prs.slides.add_slide(blank)
    panel(s, 0, 0, 13.333, 7.5, fill=NAVY, shape=MSO_SHAPE.RECTANGLE)
    panel(s, 0.8, 2.95, 1.6, 0.08, fill=RED, shape=MSO_SHAPE.RECTANGLE)
    textbox(s, 0.8, 1.2, 11.5, 1.7, ["Demo: Bảo vệ template sinh trắc học", "(Biometric Template Protection)"],
            size=36, color=WHITE, bold=True, space_after=2)
    textbox(s, 0.8, 3.2, 11.5, 0.9, ["Phần 4 · Người 4 · Nhóm 8 — MTH042 Advanced Information System Security",
                                     "HCMUS · 13/10/2026"], size=17, color=WHITE, space_after=4)
    chips = [("Tab 1–2", "So khớp và tấn công", "sau Người 1"),
             ("Tab 3", "BioHashing", "sau Người 2"),
             ("Tab 4", "BioHashing vs PolyProtect", "sau Người 3")]
    for i, (a, b, c) in enumerate(chips):
        label_box(s, 0.8 + i * 3.95, 4.75, 3.7, 1.15, f"{a}: {b}\n{c}", fill=RGBColor(0x2C, 0x4F, 0x7C), size=16)
    textbox(s, 0.8, 6.5, 11.5, 0.5, ["Ứng dụng Gradio chạy offline trên máy thuyết trình; số liệu trên slide do chính ứng dụng sinh ra."],
            size=12, color=RGBColor(0xC8, 0xD3, 0xE0))
    say = ("Phần demo gồm 4 tab, chen vào sau từng phần lý thuyết: tab 1 và 2 sau phần tấn công của Người 1, "
           "tab 3 sau BioHashing, tab 4 sau PolyProtect.")
    script.append({"slide": 1, "title": "Trang tiêu đề", "when": "Mở đầu (không cần nói nếu nhóm đã giới thiệu)",
                   "time": "0:00", "do": ["Chỉ chiếu khi bắt đầu phần demo đầu tiên."], "say": say})

    # 2. Setup / Thiết lập
    s = prs.slides.add_slide(blank)
    header(s, "Thiết lập demo (Demo setup)", "Trước Tab 1")
    y0 = 1.65
    label_box(s, 0.5, y0 + 0.75, 1.9, 0.95, "Ảnh khuôn mặt\nJPG/PNG", NAVY)
    label_box(s, 2.85, y0 + 0.75, 2.2, 0.95, "Facenet\nqua DeepFace", NAVY)
    label_box(s, 5.5, y0 + 0.75, 2.3, 0.95, "Embedding\n128 chiều, chuẩn hóa L2", NAVY)
    arrow(s, 2.4, y0 + 1.22, 2.85, y0 + 1.22)
    arrow(s, 5.05, y0 + 1.22, 5.5, y0 + 1.22)
    branches = [("Không bảo vệ\nso bằng cosine", RED), ("BioHashing\n128 bit, so bằng Hamming", BLUE),
                ("PolyProtect\nđa thức, so bằng cosine", GREEN)]
    for i, (txt, col) in enumerate(branches):
        by = y0 + i * 0.86
        label_box(s, 8.35, by, 2.45, 0.74, txt, col, size=13)
        arrow(s, 7.8, y0 + 1.22, 8.35, by + 0.37)
        arrow(s, 10.8, by + 0.37, 11.25, y0 + 1.22)
    label_box(s, 11.25, y0 + 0.75, 1.6, 0.95, "So với ngưỡng\nMATCH hoặc\nNO MATCH", NAVY, size=13)
    rows = [
        ["Thành phần", "Thiết lập trong demo", "Nguồn"],
        ["Trích đặc trưng (feature extraction)", "Facenet → vector 128 chiều, chuẩn hóa L2", "Schroff et al. 2015; Serengil & Ozpinar 2020"],
        ["Ngưỡng quyết định (threshold)", f"Cosine ≥ {th} (khoảng cách cosine 0,40)", "DeepFace, config/threshold.py"],
        ["BioHashing", "128 vector trực chuẩn sinh từ key, lấy dấu → 128 bit", "Teoh et al. 2004"],
        ["PolyProtect", f"m = 5; C ∈ [−50, 50] khác 0; E hoán vị 1..5; overlap {t4['overlap']} → {t4['template_length']} phần tử",
         "Hahn & Marcel 2022"],
        ["Tấn công leo đồi (hill-climbing)", f"{t2['iterations']} vòng, nhiễu σ = {vn(t2['sigma'], 1)}, mỗi vòng hỏi matcher 1 lần",
         "Adler 2003"],
    ]
    table(s, 0.5, 4.4, 12.35, rows, [3.3, 5.75, 3.3], size=12, row_h=0.4)
    footer(s, f"Số liệu ở các slide sau: chạy trên {n['n_images']} ảnh của {n['n_people']} người, sinh bởi scripts/export_figures.py.", 2)
    say = ("Ảnh đi qua Facenet để lấy vector 128 chiều. Từ vector này có ba cách lưu: lưu thẳng, BioHashing, "
           f"hoặc PolyProtect. Ngưỡng {th} là mặc định của thư viện DeepFace cho Facenet, không phải nhóm tự chọn.")
    script.append({"slide": 2, "title": "Thiết lập demo", "when": "Ngay sau Người 1, trước Tab 1", "time": "0:20",
                   "do": ["Chiếu slide, giới thiệu luồng xử lý trong 20 giây."], "say": say})

    # 3. Tab 1 / Tab 1
    s = prs.slides.add_slide(blank)
    header(s, "Tab 1 — Đăng ký và so khớp (Enroll & Match)", "Sau Người 1")
    picture(s, FIG / "tab1_matrix.png", 0.5, 1.4, 6.0, 5.4)
    textbox(s, 6.85, 1.5, 6.0, 3.4, [
        "Mỗi ảnh → embedding 128 số thực; so khớp bằng độ tương đồng cosine (cosine similarity).",
        f"**Cùng người:** {vn(t1['same_min'])} – {vn(t1['same_max'])} → MATCH.",
        f"**Khác người:** {vn(t1['diff_min'])} – {vn(t1['diff_max'])} → NO MATCH.",
        f"Ngưỡng {th} là điểm đánh đổi giữa tỉ lệ chấp nhận sai (FAR) và tỉ lệ từ chối sai (FRR).",
    ], size=17, bullets=True, space_after=10)
    callout(s, 6.85, 5.1, 6.0, 1.55, "Template trong CSDL là vector thô: lộ CSDL là lộ đặc trưng khuôn mặt — "
            "mà khuôn mặt thì không đổi được như mật khẩu.")
    footer(s, "Nguồn: Kim & Solomon (2018), Ch. 5, tr. 148–149 (FAR, FRR, CER); DeepFace config/threshold.py.", 3)
    say = (f"Hai ảnh cùng người cho cosine từ {vn(t1['same_min'])} đến {vn(t1['same_max'])}, vượt ngưỡng nên MATCH; "
           f"khác người chỉ quanh 0. Hệ thống chạy tốt, nhưng thứ nằm trong cơ sở dữ liệu chính là vector khuôn mặt thô. "
           "Nếu cơ sở dữ liệu bị lộ thì sao?")
    script.append({"slide": 3, "title": "Tab 1 — Đăng ký và so khớp", "when": "Sau Người 1", "time": "0:30",
                   "do": ["Chuyển sang app, Tab 1.", f"Chọn cặp ảnh cùng người trong ô Ảnh mẫu → bấm **So khớp** (kỳ vọng ≈ {vn(t1['same_min'])}–{vn(t1['same_max'])}, MATCH).",
                          "Chọn cặp khác người → bấm lại (kỳ vọng ≈ 0, NO MATCH).",
                          "Nếu app lỗi: chiếu slide này, ma trận cosine đã có đủ số."], "say": say})

    # 4. Tab 2 / Tab 2
    s = prs.slides.add_slide(blank)
    header(s, "Tab 2 — Tấn công leo đồi (Hill-climbing attack)", "Sau Người 1")
    picture(s, FIG / "tab2_attack.png", 0.4, 1.45, 6.6, 4.1)
    textbox(s, 7.25, 1.5, 5.6, 3.5, [
        "**Giả định:** kẻ tấn công chỉ gửi mẫu thử (probe) và nhận lại điểm so khớp; không thấy template.",
        "**Mỗi vòng:** thêm nhiễu nhỏ vào probe, giữ lại nếu điểm tăng.",
        f"**Kết quả {len(t2['runs'])} lần chạy:** vượt ngưỡng ở vòng {min(crossed)}–{max(crossed)}; "
        f"sau {t2['iterations']} vòng đạt {vn(min(finals))} – {vn(max(finals))}.",
    ], size=17, bullets=True, space_after=10)
    callout(s, 0.5, 5.75, 12.35, 0.95, "Điểm so khớp là kênh rò rỉ. BTP không chặn kênh này; BTP bảo vệ template khi bị lộ: "
            "không suy ra được vector gốc, và thu hồi được bằng cách đổi key.", size=15)
    footer(s, "Nguồn: Adler (2003), CCECE — hill-climbing trên template khuôn mặt dựa vào điểm so khớp.", 4)
    say = ("Kẻ tấn công không thấy template, chỉ thấy điểm. Bắt đầu từ vector ngẫu nhiên, mỗi vòng thêm chút nhiễu, "
           f"điểm tăng thì giữ. Khoảng vòng {min(crossed)} đến {max(crossed)} là vượt ngưỡng. "
           "Đây là lý do cần bảo vệ template — phần tiếp theo, Người 2 trình bày giải pháp cổ điển là BioHashing.")
    script.append({"slide": 4, "title": "Tab 2 — Tấn công leo đồi", "when": "Ngay sau Tab 1", "time": "1:00",
                   "do": ["Chuyển Tab 2, chọn một ảnh làm nạn nhân, giữ mặc định 300 vòng, σ = 0,2, seed 0.",
                          "Bấm **Chạy Attack**, im lặng khoảng 2 giây cho biểu đồ chạy.",
                          "Chỉ vào đường ngưỡng nét đứt và dòng **Attack SUCCEEDED**."], "say": say})

    # 5. Tab 3 / Tab 3
    s = prs.slides.add_slide(blank)
    header(s, "Tab 3 — BioHashing: cùng khuôn mặt, hai key", "Sau Người 2")
    picture(s, FIG / "tab3_bits.png", 0.5, 1.35, 12.35, 2.95)
    ka, kb = t3["keys"]
    fmt = lambda b: " ".join(b[i:i + 8] for i in range(0, 32, 8))
    textbox(s, 0.5, 4.5, 6.0, 0.9, [f"key = {ka}:  {fmt(t3['first32_a'])}", f"key = {kb}:  {fmt(t3['first32_b'])}"],
            size=15, font="Courier New", space_after=2)
    textbox(s, 0.5, 5.45, 6.0, 1.3, [
        f"**Cùng key:** Hamming = {vn(t3['same_key'], 3)} → MATCH.",
        f"**Khác key:** Hamming = {vn(t3['cross_key'], 3)} ≈ 0,5 → như đoán ngẫu nhiên.",
    ], size=16, bullets=True)
    textbox(s, 6.85, 4.5, 6.0, 2.3, [
        "**Tính thu hồi (renewability):** lộ template → cấp key mới → template mới không khớp template cũ.",
        "**Tính không liên kết (unlinkability):** hai hệ thống dùng hai key → không ghép được hai CSDL.",
        "**Giới hạn:** lộ cả key và template thì chỉ còn là một phép chiếu tuyến tính.",
    ], size=16, bullets=True, space_after=8)
    footer(s, "Nguồn: Teoh, Ngo & Goh (2004), Pattern Recognition 37(11), 2245–2255.", 5)
    say = (f"Cùng một khuôn mặt, key {ka} và key {kb} cho hai chuỗi bit trông hoàn toàn khác nhau. "
           f"Hamming similarity là {vn(t3['cross_key'], 3)}, đúng mức đoán ngẫu nhiên, nên key cũ đã được thu hồi an toàn. "
           "Nếu cùng key thì giống hệt, bằng 1.")
    script.append({"slide": 5, "title": "Tab 3 — BioHashing", "when": "Sau Người 2", "time": "0:45",
                   "do": [f"Chuyển Tab 3, chọn một ảnh, key_a = {ka}, key_b = {kb} → bấm **Tạo protected templates**.",
                          "Chỉ vào hai dải bit khác nhau và con số ≈ 0,5.",
                          f"Đổi key_b = {ka} → bấm lại → Hamming = 1,000."], "say": say})

    # 6. Tab 4 / Tab 4
    s = prs.slides.add_slide(blank)
    header(s, "Tab 4 — Liên kết cùng một người giữa hai CSDL", "Sau Người 3")
    picture(s, FIG / "tab4_linkage.png", 0.4, 1.4, 6.5, 3.75)
    rows = [
        ["Phương pháp", "Cùng key", f"Khác key (TB {t4['n_key_pairs']} cặp)", "Mức ngẫu nhiên"],
        ["Không bảo vệ", vn(t4["unprotected"]), f"{vn(t4['unprotected'])} (không có key)", "0"],
        ["BioHashing (128 bit)", vn(t4["bio_same_key"]), vn(t4["bio_cross_mean"]), "0,5"],
        [f"PolyProtect ({t4['template_length']} phần tử)", vn(t4["poly_same_key"]), vn(t4["poly_cross_mean"]), "0"],
    ]
    table(s, 7.15, 1.55, 5.7, rows, [2.05, 1.0, 1.65, 1.0], size=12, row_h=0.5)
    textbox(s, 7.15, 3.85, 5.7, 1.5, [
        "**Cùng key:** cả hai vẫn nhận ra đúng người, gần bằng không bảo vệ.",
        "**Khác key:** cả hai về mức ngẫu nhiên → không liên kết được hai CSDL.",
    ], size=15, bullets=True, space_after=6)
    callout(s, 0.5, 5.45, 12.35, 1.25, "0,5 trên thang Hamming và 0 trên thang cosine cùng nghĩa \"không liên quan\": "
            "đổi Hamming h sang thang tương quan bằng\u00a02h\u00a0−\u00a01 thì 0,5 ứng với 0.", fill=LIGHT, size=15)
    footer(s, f"Ảnh {pair} (cùng một người, hai lần chụp); key A = {t4['keys'][0]}, key B = {t4['keys'][1]}. "
              "Nguồn: Hahn & Marcel (2022); Teoh et al. (2004).", 6)
    say = ("Cùng một người đăng ký ở hai hệ thống dùng hai key khác nhau. Không bảo vệ thì điểm cao, ghép được ngay. "
           f"BioHashing về {vn(t4['bio_cross_mean'])}, PolyProtect về {vn(t4['poly_cross_mean'])}: cả hai là mức ngẫu nhiên, "
           f"trong khi cùng key vẫn nhận ra đúng người ({vn(t4['bio_same_key'])} và {vn(t4['poly_same_key'])}).")
    script.append({"slide": 6, "title": "Tab 4 — BioHashing vs PolyProtect", "when": "Sau Người 3", "time": "0:50",
                   "do": ["Chuyển Tab 4, chọn cặp ảnh cùng người trong ô Ảnh mẫu.",
                          f"Giữ key {t4['keys'][0]}/{t4['keys'][1]}, overlap {t4['overlap']} → bấm **So sánh 3 phương pháp**.",
                          "Chỉ biểu đồ (trung bình 200 cặp key), sau đó chỉ cột *Cùng key* trong bảng."], "say": say})

    # 7. Finding / Phát hiện
    s = prs.slides.add_slide(blank)
    header(s, "Phát hiện: không phải cặp key nào cũng an toàn", "Sau Tab 4")
    picture(s, FIG / "finding_powers.png", 0.5, 1.45, 5.4, 3.35)
    term = lambda i: [(f"c", None), (str(i), "sub"), ("v", None), (str(i), "sub"), (f"e{i}", "sup")]
    formula(s, 0.5, 4.85, 5.4, 0.6, [("p = ", None)] + term(1) + [(" + ", None)] + term(2)
            + [(" + … + ", None)] + term(5), size=20)
    k42, k99 = t4["key_42"], t4["key_99"]
    pos = k42["E"].index(1)
    same_pos = pos == k99["E"].index(1)
    cause = (f"vị trí {pos + 1} có số mũ 1 ở cả hai key, hệ số {k42['C'][pos]} và {k99['C'][pos]} cùng dấu"
             if same_pos else "hai key có hạng tử bậc 1 tương quan")
    textbox(s, 6.25, 1.5, 6.6, 3.9, [
        f"Cặp key 42/99: PolyProtect khác key vẫn cho cosine {vn(t4['poly_cross_42_99'])} → **liên kết được**.",
        f"Nguyên nhân: |v| ≤ {vn(t4['max_abs_v'])} nên hạng tử bậc 1 chi phối; {cause}.",
        f"Trên {t4['n_key_pairs']} cặp key ngẫu nhiên: **{t4['poly_linkable_rate']:.0%}** có |cosine| > 0,5.",
        "Khớp với paper: chọn (C, E) ngẫu nhiên → D↔sys 0,14–0,16; chọn chặt (strict) → 0,03–0,04 (FaceNet).",
    ], size=16, bullets=True, space_after=9)
    callout(s, 0.5, 5.6, 12.35, 1.1, "Bài học: đánh giá tính không liên kết trên nhiều cặp key, không kết luận từ một cặp; "
            "khi triển khai cần quy tắc chọn (C, E) thay vì lấy ngẫu nhiên thuần.", size=15)
    footer(s, "Nguồn: Hahn & Marcel (2022) — số D↔sys (qua Bảng 19 tài liệu nhóm); các số khác: nhóm đo trên ảnh demo.", 7)
    say = (f"Khi cài đặt, nhóm thấy cặp key 42 và 99 cho PolyProtect vẫn liên kết được, cosine {vn(t4['poly_cross_42_99'])}. "
           "Lý do: giá trị embedding rất nhỏ nên lũy thừa bậc cao gần như bằng 0, hạng tử bậc 1 quyết định; "
           "hai key này lại đặt số mũ 1 cùng chỗ. Paper cũng cho thấy chọn key theo quy tắc chặt làm điểm liên kết giảm rõ.")
    script.append({"slide": 7, "title": "Phát hiện khi cài đặt", "when": "Ngay sau Tab 4", "time": "0:35",
                   "do": ["(Tùy chọn, nếu còn thời gian) ở Tab 4 nhập key B = 99 → bấm lại → app hiện cảnh báo ⚠️."],
                   "say": say})

    # 8. Summary + references / Tổng kết + tài liệu
    s = prs.slides.add_slide(blank)
    header(s, "Tổng kết demo và tài liệu tham khảo")
    cards = [
        ("Không bảo vệ", RED, f"So khớp tốt, nhưng lộ template là lộ vĩnh viễn; leo đồi vượt ngưỡng sau ~{min(crossed)}–{max(crossed)} lần hỏi."),
        ("BioHashing", BLUE, f"Thu hồi được, không liên kết (≈ {vn(t4['bio_cross_mean'])}); an toàn phụ thuộc bí mật của key."),
        ("PolyProtect", GREEN, f"Giữ độ chính xác ({vn(t4['poly_same_key'])} so với {vn(t4['unprotected'])}), "
                               f"không liên kết trung bình (≈ {vn(t4['poly_cross_mean'])}); cần quy tắc chọn (C, E)."),
    ]
    for i, (name, col, body) in enumerate(cards):
        x = 0.5 + i * 4.2
        label_box(s, x, 1.45, 3.95, 0.55, name, col, size=16)
        panel(s, x, 2.0, 3.95, 1.55, fill=LIGHT)
        textbox(s, x + 0.12, 2.05, 3.7, 1.45, [body], size=15, anchor=MSO_ANCHOR.MIDDLE)
    textbox(s, 0.5, 3.85, 12.35, 0.4, ["Tài liệu tham khảo"], size=16, color=NAVY, bold=True)
    textbox(s, 0.5, 4.3, 12.35, 2.6, [f"[{i}] {r}" for i, r in enumerate(REFERENCES, 1)], size=12.5, space_after=4)
    say = ("Tóm lại: không bảo vệ thì lộ là mất vĩnh viễn; BioHashing cho thu hồi và không liên kết nhưng dựa vào bí mật của key; "
           "PolyProtect giữ độ chính xác và không liên kết trung bình, với điều kiện chọn key cẩn thận.")
    script.append({"slide": 8, "title": "Tổng kết", "when": "Kết thúc phần demo", "time": "0:20",
                   "do": ["Chiếu slide, nói 3 ý rồi bàn giao cho phần kết luận/Q&A."], "say": say})
    return prs, script


# ---------- Markdown script / Kịch bản markdown ----------


QA = [
    ("Vì sao Facenet trả về 128 chiều? 512 chiều có tốt hơn không?",
     "Bài báo FaceNet chọn 128 chiều; các chiều lớn hơn mà nhóm tác giả thử không cho cải thiện có ý nghĩa thống kê, "
     "trong khi template to hơn (Schroff et al., 2015). Demo dùng Facenet 128 chiều vì PolyProtect được đánh giá trên embedding 128 chiều."),
    ("Ngưỡng {th} lấy ở đâu? Doc ghi 0,5?",
     "Lấy từ cấu hình mặc định của DeepFace cho Facenet: khoảng cách cosine 0,40, tức độ tương đồng ≥ 0,60 "
     "(deepface/config/threshold.py). Tăng ngưỡng thì FAR giảm, FRR tăng (Kim & Solomon, 2018, Ch. 5). App có thanh trượt để đổi ngưỡng."),
    ("Hill-climbing hoạt động thế nào, vì sao hội tụ?",
     "Mỗi vòng thêm nhiễu nhỏ, chỉ giữ khi điểm tăng, nên điểm không bao giờ giảm. Cosine thay đổi liên tục theo vector, "
     "nên các bước nhỏ ngẫu nhiên vẫn có xác suất cải thiện đáng kể (Adler, 2003)."),
    ("Chạy 1000 vòng thay vì 300 thì sao?",
     "Điểm vẫn tăng nhưng chậm dần: trên ảnh demo, seed 0 đạt {s300} sau 300 vòng và {s1000} sau 1000 vòng."),
    ("Có BTP rồi thì hill-climbing thất bại?",
     "Không hẳn. Nếu matcher vẫn trả điểm thì kẻ tấn công vẫn leo được trong miền đã bảo vệ. BTP bảo vệ template khi bị lộ "
     "(không suy ra vector gốc, thu hồi được). Chống leo đồi cần thêm: giới hạn số lần thử, không trả điểm thô. "
     "Riêng lượng tử hóa điểm (khuyến nghị BioAPI) đã được chứng minh là chưa đủ (Adler, 2004)."),
    ("Vì sao BioHashing ra 0,5 còn PolyProtect ra 0?",
     "Hai thang đo khác nhau: BioHashing so bằng Hamming (tỉ lệ bit trùng, ngẫu nhiên = 0,5), PolyProtect so bằng cosine "
     "(ngẫu nhiên = 0). Đổi Hamming h sang thang tương quan bằng 2h − 1 thì 0,5 ứng với 0: cùng nghĩa \"không liên quan\"."),
    ("Vì sao biểu đồ Tab 4 lấy trung bình 200 cặp key?",
     "Vì một cặp key cụ thể có thể \"xui\": cặp 42/99 vẫn liên kết được (cosine {k4299}). Trên 200 cặp ngẫu nhiên, "
     "{rate} có |cosine| > 0,5. Paper cũng cho thấy chọn (C, E) chặt giảm D↔sys từ 0,14–0,16 xuống 0,03–0,04."),
    ("Overlap ảnh hưởng gì?",
     "Overlap 0→4 làm template dài 26→124 phần tử: độ chính xác tăng (TMR 96,70% → 99,87% với FaceNet) nhưng khả năng "
     "bị đảo ngược cũng tăng (0% → 95%). Demo chọn overlap 2: TMR 99,59%, tỉ lệ đảo ngược 1% (Hahn & Marcel, 2022; Bảng 18, 20 tài liệu nhóm)."),
    ("Vì sao demo chạy được khi không có mạng?",
     "Weights Facenet (~92MB) nằm sẵn trong thư mục models/ của dự án, embedding của ảnh demo được cache trong cache/; "
     "app không cần tải gì khi chạy."),
    ("Vì sao dùng Gradio?",
     "Gradio dựng giao diện web từ hàm Python trong vài dòng, chạy local trên máy thuyết trình; phần thuật toán tách riêng trong btp/ và có unit test."),
    ("Ảnh của nhóm có bị lưu lại không?",
     "Ảnh và cache embedding chỉ nằm trên máy, bị loại khỏi Git bằng .gitignore. Lưu ý trung thực: cache embedding chính là template "
     "chưa bảo vệ, nên nhóm xóa thư mục cache/ sau buổi trình bày."),
]


def write_script(n: dict, script: list[dict]) -> None:
    t2, t4 = n["tab2"], n["tab4"]
    fill = {"th": vn(n["threshold"]), "s300": vn(t2["seed0_at_300"]), "s1000": vn(t2["seed0_at_1000"]),
            "k4299": vn(t4["poly_cross_42_99"]), "rate": f"{t4['poly_linkable_rate']:.0%}"}
    lines = [
        "# Kịch bản demo — Phần 4 (Người 4)",
        "",
        f"Sinh tự động bởi `slides/build_slides.py` từ `slides/figures/numbers.json` ({n['generated_at']}, "
        f"{n['n_images']} ảnh của {n['n_people']} người). Đổi ảnh thì chạy lại hai lệnh ở cuối file.",
        "",
        "## Lịch chạy",
        "",
        "| Đoạn | Slide | Khi nào | Thời lượng |",
        "|---|---|---|---|",
        "| 1 | 2–4 | Ngay sau Người 1 | ~1:50 |",
        "| 2 | 5 | Ngay sau Người 2 | ~0:45 |",
        "| 3 | 6–8 | Ngay sau Người 3 | ~1:45 |",
        "",
        "Trước buổi trình bày: mở sẵn app (`.venv/bin/python demo_app.py`) ở một tab trình duyệt, slide ở cửa sổ khác. "
        "Nếu app lỗi, chiếu slide tương ứng: mọi biểu đồ và số liệu đã có sẵn trên slide.",
        "",
    ]
    for item in script:
        lines += [f"## Slide {item['slide']} — {item['title']}", "",
                  f"**Khi nào:** {item['when']} · **Thời lượng:** {item['time']}", "", "**Thao tác:**", ""]
        lines += [f"{i}. {step}" for i, step in enumerate(item["do"], 1)]
        lines += ["", "**Lời nói:**", "", f"> {item['say']}", ""]
    lines += ["## Chuẩn bị Q&A", ""]
    for q, a in QA:
        lines += [f"**Hỏi:** {q.format(**fill)}", "", f"**Đáp:** {a.format(**fill)}", ""]
    lines += ["## Tài liệu tham khảo", ""] + [f"{i}. {r}" for i, r in enumerate(REFERENCES, 1)]
    lines += ["", "## Cập nhật khi có ảnh của nhóm", "", "```bash",
              ".venv/bin/python scripts/export_figures.py", ".venv/bin/python slides/build_slides.py", "```", ""]
    OUT_MD.write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    numbers = json.loads((FIG / "numbers.json").read_text(encoding="utf-8"))
    prs, script = build(numbers)
    prs.save(OUT_PPTX)
    write_script(numbers, script)
    print(f"Wrote {OUT_PPTX.name} ({len(prs.slides)} slides) and {OUT_MD.name}")


if __name__ == "__main__":
    main()
